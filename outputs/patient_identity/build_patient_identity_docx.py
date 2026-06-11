from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("C:/Users/hgeec/github/symphonix-bridge-sdk")
OUT = (
    ROOT
    / "outputs"
    / "patient_identity"
    / "Symphonix_Patient_Identity_Architecture.docx"
)
MATRIX = ROOT / "tests" / "harness" / "patient_identity" / "requirements_matrix.json"
SCENARIOS = (
    ROOT / "tests" / "harness" / "patient_identity" / "patient_identity_scenarios.json"
)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_text(
    cell, text: str, bold: bool = False, color: str | None = None
) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(9)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(
    document: Document, headers: list[str], rows: list[list[str]], widths: list[float]
) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_text(cell, header, bold=True, color="FFFFFF")
        set_cell_shading(cell, "0B5D5E")
        cell.width = Inches(widths[index])
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value)
            cells[index].width = Inches(widths[index])
    document.add_paragraph()


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


matrix = json.loads(MATRIX.read_text(encoding="utf-8"))
scenarios = json.loads(SCENARIOS.read_text(encoding="utf-8"))

document = Document()
section = document.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.8)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

styles = document.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(10)
styles["Title"].font.name = "Arial"
styles["Title"].font.size = Pt(22)
styles["Title"].font.bold = True
styles["Heading 1"].font.name = "Arial"
styles["Heading 1"].font.size = Pt(16)
styles["Heading 1"].font.bold = True
styles["Heading 1"].font.color.rgb = RGBColor(11, 93, 94)
styles["Heading 2"].font.name = "Arial"
styles["Heading 2"].font.size = Pt(13)
styles["Heading 2"].font.bold = True
styles["Heading 3"].font.name = "Arial"
styles["Heading 3"].font.size = Pt(11)
styles["Heading 3"].font.bold = True

header = section.header.paragraphs[0]
header.text = "Symphonix Health | Patient Identity Closed Loop"
header.style = styles["Normal"]
header.alignment = WD_ALIGN_PARAGRAPH.LEFT
footer = section.footer.paragraphs[0]
footer.text = "Generated architecture, requirements, scenarios, and harness report"
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT

title = document.add_paragraph(style="Title")
title.add_run("Symphonix Patient Identity Closed Loop")
subtitle = document.add_paragraph()
subtitle.add_run(
    "Architecture decision, requirements, use cases, traceability matrix, and test harness"
).italic = True

metadata_rows = [
    ["Canonical identifier", "SPID - Symphonix Patient Identifier"],
    ["Identifier system", matrix["metadata"]["canonical_identifier_system"]],
    ["NHS Number role", "External alias only, never the internal master patient id"],
    ["Requirements", str(len(matrix["requirements"]))],
    ["Scenarios", str(len(scenarios["scenarios"]))],
]
add_table(document, ["Field", "Value"], metadata_rows, [2.0, 5.7])

document.add_heading("Executive Decision", level=1)
document.add_paragraph(
    "Symphonix Health should standardise BulletTrain, Shared HealthRecord, Client Registry, and all patient-data siblings on SPID. "
    "SPID has the 10-digit modulus-11 shape used by the UK NHS number format, but it is a Symphonix identifier with a Symphonix FHIR system. "
    "The Client Registry is the only issuing authority. Other services maintain local projections and fail closed when patient identity is invalid, missing, under review, restricted, or deceased."
)

document.add_heading("Architecture Summary", level=1)
add_bullets(
    document,
    [
        "Client Registry owns SPID allocation, aliases, lifecycle status, duplicate/confused-record review, merge/link, restricted status, and death status.",
        "Shared HealthRecord and siblings own clinical/domain records but key patient context by SPID.",
        "Patient creation is propagated through a transactional outbox and patient.identity.* events.",
        "Agents can detect, explain, and recommend identity corrections, but deterministic services and human approval own safety-critical state changes.",
        "Closed-loop reconciliation compares registry and subscriber projections and replays events or opens review tasks for drift.",
    ],
)

document.add_heading("NHS/PDS Lessons Applied", level=1)
nhs_rows = [
    [
        "PDS master demographics",
        "Use Client Registry as the Symphonix master patient identity service.",
    ],
    ["Duplicate records", "Merge confirmed duplicates to one SPID and retain aliases."],
    [
        "Confused records",
        "Put records under review and require NBO-style human approval before correction.",
    ],
    [
        "Sensitive flag / stop-noting",
        "Represent as restricted identity status and suppress location/contact disclosure.",
    ],
    [
        "Death status",
        "Stop active care loops and notifications without deleting the record.",
    ],
    [
        "Incorrect death status",
        "Correct through audited reinstatement or death-status correction events.",
    ],
]
add_table(document, ["NHS Pattern", "Symphonix Design Rule"], nhs_rows, [2.3, 5.4])

document.add_page_break()
document.add_heading("Core Requirements", level=1)
req_rows = [
    [
        requirement["requirement_id"],
        requirement["priority"],
        requirement["family"].replace("_", " ").title(),
        requirement["title"],
        str(len(requirement["scenario_ids"])),
    ]
    for requirement in matrix["requirements"]
]
add_table(
    document,
    ["ID", "Priority", "Family", "Requirement", "Scenario Count"],
    req_rows,
    [1.2, 0.8, 2.0, 2.9, 0.8],
)

document.add_section(WD_SECTION.NEW_PAGE)
document.add_heading("Use Case Scenarios", level=1)
scenario_rows = [
    [
        scenario["use_case_id"],
        scenario["scenario_category"],
        scenario["priority"],
        scenario["subsystem"],
        scenario["title"],
    ]
    for scenario in scenarios["scenarios"]
]
add_table(
    document,
    ["Use Case", "Type", "Priority", "Subsystem", "Title"],
    scenario_rows,
    [1.7, 0.8, 0.7, 1.5, 3.0],
)

document.add_heading("Implementation Plan", level=1)
plan_rows = [
    [
        "1",
        "Shared contract",
        "Adopt bridge_sdk.patient_identity, event schemas, and envelope patient metadata.",
        "P0",
    ],
    [
        "2",
        "BulletTrain registry and SHR",
        "Make Client Registry the SPID authority and stop local PAT-* creation in SHR.",
        "P0",
    ],
    [
        "3",
        "Sibling migration",
        "Replace internal nhs_number and PAT-* identifiers with SPID while preserving aliases.",
        "P0",
    ],
    [
        "4",
        "Closed-loop governance",
        "Add reconciliation, human approval, audit, and agent recommendation guardrails.",
        "P1",
    ],
]
add_table(
    document, ["Phase", "Theme", "Action", "Priority"], plan_rows, [0.6, 1.6, 4.5, 1.0]
)

document.add_heading("Harness", level=1)
document.add_paragraph("Fast contract suite:")
document.add_paragraph(
    "python -m pytest tests/test_patient_identity.py tests/harness/patient_identity/test_patient_identity_contract.py -q",
    style="Intense Quote",
)
document.add_paragraph("Post-migration static enforcement:")
document.add_paragraph(
    "$env:SYMPHONIX_PATIENT_IDENTITY_ENFORCE='1'; python -m pytest tests/harness/patient_identity/test_patient_identity_contract.py -q",
    style="Intense Quote",
)

document.add_heading("Sources", level=1)
source_rows = [
    ["NHS PDS", "https://digital.nhs.uk/services/personal-demographics-service"],
    [
        "PDS National Back Office",
        "https://digital.nhs.uk/services/personal-demographics-service/national-back-office-for-the-personal-demographics-service",
    ],
    [
        "PDS data quality incidents",
        "https://digital.nhs.uk/services/personal-demographics-service/national-back-office-for-the-personal-demographics-service/data-quality-incidents",
    ],
    [
        "Access restriction / S-flag",
        "https://digital.nhs.uk/services/personal-demographics-service/restricting-access-to-a-patients-demographic-record",
    ],
    [
        "NHS number format",
        "https://digital.nhs.uk/services/personal-demographics-service/nhs-number",
    ],
    [
        "NHS Data Dictionary modulus-11",
        "https://v2.datadictionary.nhs.uk/data_dictionary/data_field_notes/n/nhs_number_de.asp%40shownav%3D0.html",
    ],
    [
        "NHS Notify death/flag suppression",
        "https://webcms.prod.nhsnotify.national.nhs.uk/using-nhs-notify/recipients-you-cannot-message",
    ],
    [
        "Records Management Code",
        "https://transform.england.nhs.uk/information-governance/guidance/records-management-code/records-management-code-of-practice/",
    ],
]
add_table(document, ["Source", "URL"], source_rows, [2.2, 5.5])

OUT.parent.mkdir(parents=True, exist_ok=True)
document.save(OUT)
print(OUT)
